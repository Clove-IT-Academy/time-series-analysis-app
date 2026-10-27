"""Bounded in-memory CSV/XLSX reading. Excel formulas are never evaluated."""

import csv
import hashlib
import io
import zipfile
from pathlib import PurePath

import pandas as pd
from openpyxl import load_workbook

from .config import MAX_BYTES, MAX_COLUMNS, MAX_ROWS, MAX_SHEETS, MAX_ZIP_BYTES
from .contracts import InputError, UploadResult


def _check(data: bytes, filename: str) -> str:
    if not data or len(data) > MAX_BYTES:
        raise InputError("Choose a nonempty file no larger than 10 MiB.")
    ext = PurePath(filename).suffix.lower()
    if ext not in (".csv", ".xlsx"):
        raise InputError("Supported formats: CSV and XLSX. Convert legacy XLS to XLSX first.")
    return ext


def _workbook(data: bytes):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if (
                len(archive.infolist()) > 2000
                or sum(i.file_size for i in archive.infolist()) > MAX_ZIP_BYTES
            ):
                raise InputError("Workbook expanded size exceeds the prototype's safe limit.")
            if any("vbaProject" in n or "externalLink" in n for n in archive.namelist()):
                raise InputError("Workbooks with macros or external links are unsupported.")
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        if len(book.sheetnames) > MAX_SHEETS:
            book.close()
            raise InputError("Workbook has more than 20 worksheets.")
        return book
    except InputError:
        raise
    except Exception:
        raise InputError("Cannot read workbook. Use an unencrypted, valid XLSX file.") from None


def list_sheets(data: bytes, filename: str) -> list[str]:
    if _check(data, filename) == ".csv":
        return []
    book = _workbook(data)
    try:
        return book.sheetnames
    finally:
        book.close()


def read_upload(data: bytes, filename: str, sheet: str | None = None) -> UploadResult:
    ext = _check(data, filename)
    try:
        if ext == ".csv":
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                raise InputError("CSV must use UTF-8 encoding. Re-save it as UTF-8 CSV.") from None
            # csv.reader bounds rows/columns before constructing a full DataFrame.
            reader = csv.reader(io.StringIO(text), strict=True)
            headers = next(reader, [])
            rows = []
            for row in reader:
                if not row:
                    continue
                if len(rows) >= MAX_ROWS or len(row) > MAX_COLUMNS:
                    raise InputError("Upload exceeds 50,000 rows or 100 columns.")
                if len(row) != len(headers):
                    raise InputError("CSV rows must have the same number of fields as the header.")
                rows.append(row)
        else:
            book = _workbook(data)
            try:
                if sheet is None:
                    if len(book.sheetnames) != 1:
                        raise InputError("Select a worksheet before reading this workbook.")
                    sheet = book.sheetnames[0]
                if sheet not in book.sheetnames:
                    raise InputError("Selected worksheet does not exist.")
                ws = book[sheet]
                # Ignore misleading dimension metadata; bounded iteration still enforces limits.
                ws.reset_dimensions()
                iterator = ws.iter_rows()
                first = next(iterator, ())
                if any(c.data_type == "f" for c in first):
                    raise InputError("Formula headers are unsupported; paste values first.")
                headers = [c.value for c in first]
                rows = []
                for cells in iterator:
                    if any(c.data_type == "f" for c in cells):
                        raise InputError(
                            "Formula cells are unsupported. Paste values into a copy first."
                        )
                    row = [c.value for c in cells]
                    if not any(v is not None for v in row):
                        continue
                    if len(rows) >= MAX_ROWS or len(row) > MAX_COLUMNS:
                        raise InputError("Upload exceeds 50,000 rows or 100 columns.")
                    rows.append(row + [None] * max(0, len(headers) - len(row)))
            finally:
                book.close()
        if not headers or len(headers) > MAX_COLUMNS or not rows:
            raise InputError("File needs a header and at least one row, with at most 100 columns.")
        if any(h is None or not str(h).strip() for h in headers):
            raise InputError("Every column needs a nonempty header.")
        headers = [str(h).strip() for h in headers]
        if len(set(headers)) != len(headers):
            raise InputError("Column headers must be unique after trimming whitespace.")
        if any(len(row) > len(headers) for row in rows):
            raise InputError("Worksheet values extend beyond the header. Add missing column names.")
        table = pd.DataFrame(rows, columns=headers)
        return UploadResult(table, PurePath(filename).name, hashlib.sha256(data).hexdigest(), sheet)
    except InputError:
        raise
    except Exception:
        raise InputError(
            "Unable to parse file. Check delimiters, headers and workbook contents."
        ) from None
